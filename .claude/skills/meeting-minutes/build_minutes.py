#!/usr/bin/env python3
"""회의 내용 JSON -> 회의록 + 보고서 엑셀(A4 1장씩)을 날짜_주제 회의 폴더에 저장.

사용: python3 build_minutes.py 회의.json [--out-root 저장위치]
저장 위치: <out-root>/YYMMDD_<주제> 회의/YYMMDD_회의록_<주제>.xlsx, YYMMDD_보고서_<주제>.xlsx
"""
import argparse, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

F = "맑은 고딕"
thin = Side(style="thin", color="888888")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill("solid", fgColor="F2F2F2")


class Sheet:
    def __init__(self, widths):
        self.wb = Workbook(); self.ws = self.wb.active
        for c, w in zip("ABCDE", widths): self.ws.column_dimensions[c].width = w
        self.r = 1; self.ncol = len(widths)

    def cell(self, col, v, bold=False, fill=None, align="left", border=True, size=10):
        c = self.ws.cell(self.r, col, v); c.font = Font(name=F, size=size, bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal=align)
        if fill: c.fill = fill
        if border: c.border = BD

    def merge(self, c1, c2, v, **k):
        for c in range(c1, c2 + 1): self.cell(c, v if c == c1 else None, **k)
        if c2 > c1: self.ws.merge_cells(start_row=self.r, start_column=c1, end_row=self.r, end_column=c2)

    def height(self, h): self.ws.row_dimensions[self.r].height = h

    def title(self, t):
        self.merge(1, self.ncol, t, bold=True, border=False, align="center", size=16); self.height(32); self.r += 2

    def sec(self, t):
        self.merge(1, self.ncol, t, bold=True, border=False, size=11)
        for c in range(1, self.ncol + 1): self.ws.cell(self.r, c).border = Border(bottom=Side(style="medium", color="000000"))
        self.height(20); self.r += 1

    def lines(self, items, per=22):
        self.merge(1, self.ncol, "\n".join("· " + s for s in items), border=False)
        self.height(per * max(len(items), 1) + 8); self.r += 2

    def finish(self, path):
        ws = self.ws
        ws.page_setup.paperSize = 9; ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = f"A1:{'ABCDE'[self.ncol-1]}{self.r}"; ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.5; ws.page_margins.top = ws.page_margins.bottom = 0.6
        self.wb.save(path)


def make_minutes(d, path):
    s = Sheet([16, 44, 30, 22]); ws = s.ws
    s.title("회의록  |  " + d["title"])
    s.cell(1, "작성일", True, HF, "center"); s.cell(2, f'{d["date"]} (장소: {d.get("place") or "확인 필요"})')
    s.cell(3, "작성자 / 참석자", True, HF, "center"); s.cell(4, f'{d.get("writer","")} / {d.get("attendees") or "확인 필요"}'); s.height(36); s.r += 2
    s.sec("1. 결론 (요약)"); s.lines(d["summary"])
    s.sec("2. 논의 내용")
    for i, h in enumerate(["구분", "내용", "비고", ""]): s.cell(i + 1, h, True, HF, "center")
    ws.merge_cells(start_row=s.r, start_column=3, end_row=s.r, end_column=4); s.height(20); s.r += 1
    for a, b, c in d["items"]:
        s.cell(1, a, True, align="center"); s.cell(2, b); s.cell(3, c); s.cell(4, None)
        ws.merge_cells(start_row=s.r, start_column=3, end_row=s.r, end_column=4)
        s.height(max(34, 15 * (len(b) // 30 + 1))); s.r += 1
    s.r += 1
    s.sec("3. 후속 조치")
    for i, h in enumerate(["No.", "할 일", "담당", "기한"]): s.cell(i + 1, h, True, HF, "center")
    s.height(20); s.r += 1
    for n, (t, w, dd) in enumerate(d["actions"], 1):
        s.cell(1, n, align="center"); s.cell(2, t); s.cell(3, w or "확인 필요"); s.cell(4, dd or "확인 필요"); s.height(32); s.r += 1
    s.r += 1
    s.sec("4. 확인 필요 (녹취 변환 오류 가능)"); s.lines(d["unclear"], 20)
    s.finish(path)


def make_report(d, path):
    R = d["report"]; s = Sheet([14, 22, 22, 22, 22]); ws = s.ws
    s.title(R["title"])
    s.merge(1, 5, f'작성일: {d["date"]}   |   {R.get("dept","")} / {d.get("writer","")}   |   수신: {R.get("to","")}', border=False, align="right")
    s.height(22); s.r += 2

    def rows(items):
        for a, b in items:
            s.cell(1, a, True, HF, "center"); s.merge(2, 5, b); s.height(max(26, 15 * (len(b) // 46 + 1) + 6)); s.r += 1
        s.r += 1
    s.sec("1. 개요"); rows(R["overview"])
    s.sec("2. 주요 내용"); rows(R["main"])
    s.sec("3. 검토 의견 (안)"); rows(R["opinion"])
    s.sec("4. 향후 계획 (안)")
    s.cell(1, "No.", True, HF, "center"); s.merge(2, 3, "할 일", bold=True, fill=HF, align="center")
    s.cell(4, "담당", True, HF, "center"); s.cell(5, "기한", True, HF, "center"); s.r += 1
    for n, (t, w, dd) in enumerate(R["plan"], 1):
        s.cell(1, n, align="center"); s.merge(2, 3, t); s.cell(4, w or "확인 필요"); s.cell(5, dd or "확인 필요"); s.height(32); s.r += 1
    s.r += 1
    s.sec("5. 확인 필요"); s.lines(d["unclear"], 20)
    s.finish(path)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("json_path"); ap.add_argument("--out-root", default="회의록")
    a = ap.parse_args()
    d = json.load(open(a.json_path, encoding="utf-8"))
    ymd = d["date"].replace("-", "")[2:]
    folder = os.path.join(a.out_root, f'{ymd}_{d["topic"]} 회의')
    os.makedirs(folder, exist_ok=True)
    p1 = os.path.join(folder, f'{ymd}_회의록_{d["topic"]}.xlsx'); p2 = os.path.join(folder, f'{ymd}_보고서_{d["topic"]}.xlsx')
    make_minutes(d, p1); make_report(d, p2)
    print(p1); print(p2)


if __name__ == "__main__":
    main()
