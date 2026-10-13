# -*- coding: utf-8 -*-
"""거래처 로그인 런처
- 같은 폴더의 '거래처 계정 엑셀(.xlsx)'을 읽어 업체별 버튼을 만든다.
- 버튼을 누르면 로그인 페이지가 열리고 아이디/비번이 자동 입력된다.
- 아이디/비번은 엑셀에만 있고, 이 프로그램 코드/저장소에는 들어 있지 않다.
"""
import glob
import os
import sys
import threading
import tkinter as tk
from tkinter import messagebox, ttk

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))

# 계정표의 거래처명과 링크표의 거래처명이 다를 때 연결해 주는 표
LINK_ALIAS = {"사조 푸디스트": "한화-푸디스트"}

# 로그인 입력칸 찾기: 보이는 비밀번호 칸을 찾고, 그 앞에 있는 가장 가까운 글자 입력칸을 아이디로 본다.
FIND_FIELDS_JS = r"""
function vis(e){const r=e.getBoundingClientRect();const s=getComputedStyle(e);
  return r.width>0&&r.height>0&&s.visibility!=='hidden'&&s.display!=='none'&&!e.disabled;}
const all=[...document.querySelectorAll('input')].filter(vis);
const pw=all.find(e=>e.type==='password');
if(!pw) return null;
const idx=all.indexOf(pw);
const okType=t=>['text','email','tel','number','search',''].includes(t);
let id=null;
for(let i=idx-1;i>=0;i--){ if(okType(all[i].type)){id=all[i];break;} }
if(!id) return null;
return [id,pw];
"""


def load_accounts(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    links = {}
    for r in range(2, ws.max_row + 1):
        name, url = ws.cell(r, 8).value, ws.cell(r, 9).value
        if name and url and str(url).startswith("http"):
            links[str(name).strip()] = str(url).strip()
    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        uid, pw = ws.cell(r, 5).value, ws.cell(r, 6).value
        if not name or uid is None or pw is None:
            continue
        name = str(name).strip()
        url = links.get(name) or links.get(LINK_ALIAS.get(name, ""))
        rows.append(dict(
            name=name, manager=str(ws.cell(r, 2).value or ""),
            mode=str(ws.cell(r, 3).value or ""), uid=str(uid).strip(),
            pw=str(pw).strip(), url=url))
    return rows


def find_excel():
    if len(sys.argv) > 1:
        return sys.argv[1]
    for p in glob.glob(os.path.join(BASE, "*.xlsx")):
        if os.path.basename(p).startswith("~$"):
            continue
        try:
            if openpyxl.load_workbook(p, read_only=True).active["A1"].value == "거래처":
                return p
        except Exception:
            pass
    return None


def new_driver():
    from selenium import webdriver
    last = None
    for kind in ("chrome", "edge"):
        try:
            if kind == "chrome":
                o = webdriver.ChromeOptions()
                o.add_experimental_option("detach", True)  # 프로그램을 닫아도 브라우저는 유지
                return webdriver.Chrome(options=o)
            o = webdriver.EdgeOptions()
            o.add_experimental_option("detach", True)
            return webdriver.Edge(options=o)
        except Exception as e:  # 크롬이 없으면 엣지로
            last = e
    raise last


def fill_login(driver, uid, pw, submit=True, timeout=20):
    """로그인 칸이 나타날 때까지 기다렸다가(메인 페이지 + iframe) 입력. 성공 여부를 돌려준다."""
    from selenium.webdriver.common.keys import Keys
    import time
    end = time.time() + timeout
    while time.time() < end:
        for frame in [None] + driver.find_elements("css selector", "iframe"):
            try:
                driver.switch_to.default_content()
                if frame is not None:
                    driver.switch_to.frame(frame)
                found = driver.execute_script("return (function(){" + FIND_FIELDS_JS + "})()")
            except Exception:
                continue
            if found:
                idf, pwf = found
                idf.clear(); idf.send_keys(uid)
                pwf.clear(); pwf.send_keys(pw)
                if submit:
                    pwf.send_keys(Keys.ENTER)
                return True
        time.sleep(0.5)
    return False


class App(tk.Tk):
    def __init__(self, rows, path):
        super().__init__()
        self.title("거래처 로그인 런처")
        try:
            self.iconbitmap(os.path.join(BASE, "icon.ico"))
        except Exception:
            pass  # 아이콘 파일이 없거나 윈도우가 아니면 기본 아이콘
        self.rows = rows
        self.auto_submit = tk.BooleanVar(value=True)
        self.status = tk.StringVar(value=f"계정표: {os.path.basename(path)}")
        top = ttk.Frame(self, padding=8); top.pack(fill="x")
        ttk.Checkbutton(top, text="입력 후 자동으로 로그인 누르기(Enter)", variable=self.auto_submit).pack(side="left")
        body = ttk.Frame(self, padding=8); body.pack(fill="both", expand=True)
        for i, h in enumerate(["거래처", "접속방법", "", "", ""]):
            ttk.Label(body, text=h, font=("", 9, "bold")).grid(row=0, column=i, padx=4, sticky="w")
        for n, r in enumerate(rows, 1):
            ttk.Button(body, text=r["name"], width=18,
                       command=lambda r=r: self.login(r)).grid(row=n, column=0, padx=4, pady=2, sticky="w")
            ttk.Label(body, text=r["mode"]).grid(row=n, column=1, padx=4, sticky="w")
            if r["uid"]:
                ttk.Button(body, text="ID복사", width=7,
                           command=lambda r=r: self.copy(r["uid"], r["name"] + " 아이디")).grid(row=n, column=2)
                ttk.Button(body, text="비번복사", width=8,
                           command=lambda r=r: self.copy(r["pw"], r["name"] + " 비밀번호")).grid(row=n, column=3)
        ttk.Label(self, textvariable=self.status, padding=8, foreground="#444").pack(fill="x")

    def copy(self, text, what):
        self.clipboard_clear(); self.clipboard_append(text)
        self.status.set(f"{what} 복사됨 (붙여넣기 하세요)")

    def login(self, r):
        if not r["url"]:
            return messagebox.showwarning("링크 없음", f"{r['name']}의 로그인 링크가 엑셀(H·I열)에 없습니다.")
        if r["mode"] in ("IE호환모드", "별도툴"):
            # 자동 입력이 안 되는 방식: 링크만 열고 아이디를 복사해 둔다.
            import webbrowser
            webbrowser.open(r["url"])
            if r["uid"]:
                self.copy(r["uid"], r["name"] + " 아이디")
            self.status.set(f"{r['name']}: {r['mode']} 사이트라 직접 로그인 필요. 아이디 복사됨, 비번은 '비번복사'")
            return
        if not r["uid"]:
            import webbrowser; webbrowser.open(r["url"]); return
        self.status.set(f"{r['name']} 여는 중...")
        threading.Thread(target=self._run, args=(r,), daemon=True).start()

    def _run(self, r):
        try:
            d = new_driver()
            d.get(r["url"])
            ok = fill_login(d, r["uid"], r["pw"], self.auto_submit.get())
            msg = "입력 완료" if ok else "로그인 칸을 못 찾음 → 'ID복사/비번복사'로 직접 입력"
        except Exception as e:
            msg = f"실패: {str(e)[:80]}"
        self.after(0, lambda: self.status.set(f"{r['name']}: {msg}"))


if __name__ == "__main__":
    path = find_excel()
    if not path:
        tk.Tk().withdraw()
        messagebox.showerror("계정표 없음", "이 폴더에 거래처 계정 엑셀(.xlsx)을 넣어 주세요.\n(A1 칸이 '거래처'인 파일)")
        sys.exit(1)
    App(load_accounts(path), path).mainloop()
